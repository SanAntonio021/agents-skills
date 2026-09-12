"""Offline fault and evidence-binding tests; fixture render files are synthetic."""
import copy
import json
import os
from pathlib import Path
import sys
import subprocess
import unittest
from unittest import mock
from zipfile import ZipFile, ZIP_DEFLATED

from pptx import Presentation
from pptx.util import Pt

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import template_acceptance as acceptance
import template_deck
import render_deck
import test_template_deck as fixtures


class TemplateAcceptanceTests(unittest.TestCase):
    setUp = fixtures.TemplateDeckTests.setUp
    digest = staticmethod(fixtures.TemplateDeckTests.digest)

    def prepare(self):
        self.spec_path.write_text(json.dumps(self.spec), encoding="utf-8")
        self.prepared = self.root / "adapted.pptx"
        self.prepared_spec = self.root / "adapted.json"
        return acceptance.prepare_template(self.spec_path, self.prepared, self.prepared_spec)

    def sample(self):
        self.prepare()
        template_deck.make_template_pptx(self.deck, render_deck.validate_deck(self.deck, self.root),
                                        self.output, self.prepared_spec, template_test=True)
        # No real renderer is used here; these bytes only test binding and failures.
        pdf = self.root / "sample.pdf"
        pdf.write_bytes(b"synthetic-pdf-evidence")
        png = self.root / "sample.png"
        png.write_bytes(b"synthetic-png-evidence")
        self.manifest = self.root / "manifest.json"
        report = {"schema_version": 3, "template_test": True, "slide_count": 1,
                  "stages": {key: "passed" for key in ("pptx", "structure", "libreoffice", "png")},
                  "template": {"sha256": self.digest(self.prepared), "spec_sha256": self.digest(self.prepared_spec)},
                  "pptx_sha256": self.digest(self.output),
                  "files": {"pptx": str(self.output), "pdf": str(pdf), "png": [str(png)]},
                  "file_hashes": {str(p): self.digest(p) for p in (self.output, pdf, png)}}
        self.manifest.write_text(json.dumps(report), encoding="utf-8")
        return report

    def test_unprepared_or_unaccepted_official_output_is_blocked(self):
        self.spec_path.write_text(json.dumps(self.spec), encoding="utf-8")
        slides = render_deck.validate_deck(self.deck, self.root)
        with self.assertRaisesRegex(ValueError, "not prepared"):
            template_deck.make_template_pptx(self.deck, slides, self.output, self.spec_path)
        self.prepare()
        with self.assertRaisesRegex(ValueError, "acceptance missing"):
            template_deck.make_template_pptx(self.deck, slides, self.output, self.prepared_spec)
        self.assertFalse(self.output.exists())

    def test_preparation_resolves_paragraph_default_without_touching_original(self):
        prs = Presentation(self.source)
        title = next(s for s in prs.slides[0].shapes if s.shape_id == self.spec["layouts"]["result"]["slots"]["title"])
        title.text_frame.paragraphs[0].runs[0].font.size = None
        title.text_frame.paragraphs[0].font.size = Pt(32)
        prs.save(self.source)
        before = self.digest(self.source)
        result = self.prepare()
        self.assertEqual(result["titles"][0]["size_pt"], 32)
        self.assertEqual(self.digest(self.source), before)
        adapted_title = next(s for s in Presentation(self.prepared).slides[0].shapes if s.shape_id == title.shape_id)
        self.assertEqual(adapted_title.text_frame.paragraphs[0].runs[0].font.size.pt, 32)

    def test_preparation_resolves_layout_title_and_master_title_style(self):
        for from_layout in (True, False):
            with self.subTest(from_layout=from_layout):
                prs = Presentation()
                slide = prs.slides.add_slide(prs.slide_layouts[0])
                slide.shapes.title.text = "Inherited title"
                if from_layout:
                    prs.slide_layouts[0].placeholders[0].text_frame.paragraphs[0].font.size = Pt(30)
                prs.save(self.source)
                with ZipFile(self.source) as archive:
                    parts = {name: archive.read(name) for name in archive.namelist()}
                _, ordered = template_deck.slide_parts(parts)
                shape = template_deck.xml(parts[ordered[0]]).find("p:cSld/p:spTree/p:sp", template_deck.NS)
                value = acceptance.effective_title_size(parts, ordered[0], shape)
                self.assertEqual(value, 3000 if from_layout else 4400)

    def test_unresolved_title_reports_object_and_retries_with_actual_user_size(self):
        with ZipFile(self.source) as archive:
            parts = {name: archive.read(name) for name in archive.namelist()}
        for name in parts:
            if name.endswith(".xml"):
                root = template_deck.xml(parts[name])
                for element in root.iter():
                    element.attrib.pop("sz", None)
                parts[name] = template_deck.dump(root)
        with ZipFile(self.source, "w", ZIP_DEFLATED) as archive:
            for name, value in parts.items():
                archive.writestr(name, value)
        original_hash = self.digest(self.source)
        with self.assertRaisesRegex(ValueError, "slide=1, shape_id=.*actual --title-size-pt"):
            self.prepare()
        self.assertFalse(self.prepared.exists())
        result = acceptance.prepare_template(self.spec_path, self.prepared, self.prepared_spec, 34)
        self.assertEqual(result["titles"][0]["size_pt"], 34)
        self.assertEqual(self.digest(self.source), original_hash)

    def test_explicit_actual_size_resolves_ambiguous_inheritance(self):
        with mock.patch.object(acceptance, "effective_title_size", side_effect=ValueError("Ambiguous title placeholder inheritance; specify the actual title size")):
            with self.assertRaisesRegex(ValueError, "slide=1, shape_id=.*Ambiguous"):
                self.prepare()
            result = acceptance.prepare_template(self.spec_path, self.prepared, self.prepared_spec, 33)
        self.assertEqual(result["titles"][0]["size_pt"], 33)
        self.assertEqual(result["titles"][0]["source"], "user_actual_size")

    def test_interrupted_preparation_reports_partial_new_output_and_preserves_original(self):
        self.spec_path.write_text(json.dumps(self.spec), encoding="utf-8")
        output = self.root / "partial.pptx"
        mapping = self.root / "missing-directory" / "mapping.json"
        with self.assertRaisesRegex(OSError, "incomplete new outputs"):
            acceptance.prepare_template(self.spec_path, output, mapping)
        self.assertTrue(output.exists())
        self.assertFalse(mapping.exists())
        self.assertEqual(self.digest(self.source), self.source_hash)

    def test_acceptance_supports_top_level_pptx_hash(self):
        report = self.sample()
        report["file_hashes"].pop(str(self.output))
        self.manifest.write_text(json.dumps(report), encoding="utf-8")
        acceptance.record_acceptance(self.prepared_spec, self.manifest, visual_inspected=True)

    def test_template_cli_emits_utf8_json_under_gbk(self):
        prs = Presentation(self.source)
        title_id = self.spec["layouts"]["result"]["slots"]["title"]
        next(s for s in prs.slides[0].shapes if s.shape_id == title_id).text = "测试标题🙂"
        prs.save(self.source)
        command = [sys.executable, str(Path(template_deck.__file__)), "inspect", "--template", str(self.source), "--out", str(self.root / "检查.json")]
        result = subprocess.run(command, capture_output=True, env={**os.environ, "PYTHONIOENCODING": "gbk"})
        self.assertEqual(result.returncode, 0, result.stderr)
        report = json.loads(result.stdout.decode("utf-8"))
        self.assertTrue(any(s["text"] == "测试标题🙂" for s in report["slides"][0]["shapes"]))
        command[command.index("--template") + 1] = str(self.root / "不存在🙂.pptx")
        failed = subprocess.run(command, capture_output=True, env={**os.environ, "PYTHONIOENCODING": "gbk"})
        self.assertEqual(failed.returncode, 1)
        self.assertFalse(json.loads(failed.stdout.decode("utf-8"))["ok"])

    def test_sample_marker_and_accepted_formal_output(self):
        self.sample()
        report = template_deck.inspect_template(self.output)
        self.assertTrue(any(s["text"] == acceptance.MARKER for s in report["slides"][0]["shapes"]))
        acceptance.record_acceptance(self.prepared_spec, self.manifest, visual_inspected=True)
        official = self.root / "official.pptx"
        template_deck.make_template_pptx(self.deck, render_deck.validate_deck(self.deck, self.root), official, self.prepared_spec)
        self.assertFalse(any(s["text"] == acceptance.MARKER for s in template_deck.inspect_template(official)["slides"][0]["shapes"]))
        self.assertEqual(self.digest(self.source), self.source_hash)

    def test_acceptance_needs_actual_visual_confirmation_and_passed_render_stages(self):
        report = self.sample()
        with self.assertRaisesRegex(ValueError, "visually inspected"):
            acceptance.record_acceptance(self.prepared_spec, self.manifest)
        report["stages"]["png"] = "failed"
        self.manifest.write_text(json.dumps(report), encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "rendering must pass"):
            acceptance.record_acceptance(self.prepared_spec, self.manifest, visual_inspected=True)

    def test_changed_rendered_image_cannot_be_accepted(self):
        report = self.sample()
        Path(report["files"]["png"][0]).write_bytes(b"changed")
        with self.assertRaisesRegex(ValueError, "hash mismatch"):
            acceptance.record_acceptance(self.prepared_spec, self.manifest, visual_inspected=True)

    def test_receipt_stales_on_original_copy_or_mapping_change(self):
        self.sample()
        acceptance.record_acceptance(self.prepared_spec, self.manifest, visual_inspected=True)
        spec = json.loads(self.prepared_spec.read_text(encoding="utf-8"))
        for path in (self.source, self.prepared):
            original = path.read_bytes()
            path.write_bytes(original + b"changed")
            with self.assertRaisesRegex(ValueError, "changed after preparation"):
                acceptance.require_acceptance(spec, self.prepared_spec)
            path.write_bytes(original)
        changed = copy.deepcopy(spec)
        changed["layouts"]["result"]["keep_shape_ids"] = []
        with self.assertRaisesRegex(ValueError, "stale"):
            acceptance.require_acceptance(changed, self.prepared_spec)
        changed = copy.deepcopy(spec)
        changed["adaptation_receipt"] = str(self.prepared_spec.with_suffix(".adaptation.json"))
        self.assertEqual(acceptance.spec_digest(changed), acceptance.spec_digest(spec))
        acceptance.require_acceptance(changed, self.prepared_spec)

    def test_changed_sample_invalidates_existing_acceptance(self):
        self.sample()
        acceptance.record_acceptance(self.prepared_spec, self.manifest, visual_inspected=True)
        self.output.write_bytes(self.output.read_bytes() + b"changed")
        spec = json.loads(self.prepared_spec.read_text(encoding="utf-8"))
        with self.assertRaisesRegex(ValueError, "evidence changed"):
            acceptance.require_acceptance(spec, self.prepared_spec)

    def test_partial_layout_sample_does_not_accept_whole_mapping(self):
        self.spec["layouts"]["other"] = copy.deepcopy(self.spec["layouts"]["result"])
        self.sample()
        with self.assertRaisesRegex(ValueError, "every configured layout"):
            acceptance.record_acceptance(self.prepared_spec, self.manifest, visual_inspected=True)


if __name__ == "__main__":
    unittest.main()
