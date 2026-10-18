"""Task-local use of the unchanged bundled generator. All output stays here."""
import os
import sys
import json
import hashlib
import copy
from pathlib import Path
from datetime import datetime, timezone

sys.dont_write_bytecode = True
ROOT = Path(__file__).resolve().parent
SKILLS = ROOT.parents[2]
os.environ['TEMP'] = str(ROOT)
os.environ['TMP'] = str(ROOT)
os.environ['PYTHONDONTWRITEBYTECODE'] = '1'
sys.path.insert(0, str(SKILLS / 'lab-report-slides/scripts'))
from render_deck import validate_deck, make_pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from PIL import ImageFont

def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()

def main():
    deck_path = ROOT / 'deck.json'
    deck = json.loads(deck_path.read_text(encoding='utf-8'))
    paths = [ROOT.parent / 'input/switch.png', ROOT.parent / 'input/phase-noise.png',
             SKILLS / 'lab-report-slides/SKILL.md', SKILLS / 'pptx/SKILL.md',
             SKILLS / 'lab-report-slides/references/template-profile.json',
             SKILLS / 'lab-report-slides/scripts/render_deck.py']
    before = {str(p): sha(p) for p in paths}
    slides = validate_deck(copy.deepcopy(deck), ROOT)
    candidate = ROOT / 'builtin-generated-v03.pptx'
    assets = make_pptx(deck, slides, candidate)
    final = ROOT / '20261014_日报样例_技能流程测试_v03.pptx'
    if final.exists():
        raise FileExistsError(final)
    prs = Presentation(candidate)
    for slide in prs.slides:
        for shape in list(slide.shapes):
            if shape.has_text_frame and shape.text.strip():
                if shape.top < Inches(.75):
                    shape.name = 'Title'
                    runs = shape.text_frame.paragraphs[0].runs
                    if len(runs) == 2:
                        # Preserve the actual three ASCII spaces on the topic
                        # run and its 28 pt formatting; give it a separate role.
                        main_font = ImageFont.truetype('C:/Windows/Fonts/msyhbd.ttc', 360)
                        sub_font = ImageFont.truetype('C:/Windows/Fonts/msyh.ttc', 280)
                        main_width = Pt(main_font.getlength(runs[0].text) / 10)
                        baseline_shift = Pt((main_font.getmetrics()[0] - sub_font.getmetrics()[0]) / 10)
                        topic_xml = copy.deepcopy(shape._element)
                        props = topic_xml.find('.//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr')
                        props.set('id', str(max(s.shape_id for s in slide.shapes) + 1))
                        props.set('name', 'SectionHeader')
                        slide.shapes._spTree.insert_element_before(topic_xml, 'p:extLst')
                        topic = next(s for s in slide.shapes if s._element is topic_xml)
                        first = topic.text_frame.paragraphs[0].runs[0]._r
                        first.getparent().remove(first)
                        topic.left = shape.left + main_width
                        topic.top = shape.top + baseline_shift
                        topic.width = shape.width - main_width
                        last = runs[1]._r
                        last.getparent().remove(last)
                        shape.width = main_width + Pt(1)
                elif '方案连接图' in shape.text or '2026-10-12 仿真' in shape.text:
                    shape.name = 'Source'
        marker = slide.shapes.add_textbox(Inches(.35), Inches(7.12), Inches(5), Inches(.22))
        marker.name = 'Footnote'
        marker.text_frame.margin_left = marker.text_frame.margin_right = 0
        marker.text_frame.margin_top = marker.text_frame.margin_bottom = 0
        run = marker.text_frame.paragraphs[0].add_run()
        run.text = '技能流程测试｜日报样例'
        run.font.name = '微软雅黑'
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor.from_string('6B7280')
    prs.save(final)
    after = {str(p): sha(p) for p in paths}
    assert before == after, 'Inputs or generator changed'
    record = {
        'created_at_utc': datetime.now(timezone.utc).isoformat(),
        'authoring': 'unchanged lab-report-slides make_pptx, bundled style profile',
        'source_deck': str(deck_path), 'source_deck_sha256': sha(deck_path),
        'source_hashes_before': before, 'source_hashes_after': after,
        'source_integrity': 'PASS', 'assets': assets,
        'generated_pptx': str(candidate), 'final_pptx': str(final),
        'final_sha256': sha(final), 'slide_count': len(prs.slides),
        'editability': 'Native text and shape objects; two independently replaceable raster images.',
        'libreoffice': 'NOT_RUN: fixed external lock writes outside allowed production directory',
        'powerpoint': 'PENDING', 'visual': 'PENDING',
        'publication': 'isolated sample only; no formal-directory delivery'
    }
    (ROOT / 'production-manifest.json').write_text(json.dumps(record, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(record, ensure_ascii=False, indent=2))

if __name__ == '__main__':
    main()
