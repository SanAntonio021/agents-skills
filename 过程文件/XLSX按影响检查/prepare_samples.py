from pathlib import Path
import json
import sys
import hashlib
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.comments import Comment
from openpyxl.worksheet.datavalidation import DataValidation

root = Path(__file__).resolve().parent
fixtures = root / 'fixtures'
fixtures.mkdir(exist_ok=False)
scripts = Path(r'D:\BaiduSyncdisk\.agents\skills\xlsx\scripts')
sys.path.insert(0, str(scripts))
from libreoffice_headless import convert, find_soffice
from merge_formula_caches import merge_caches
from verify_xlsx import inspect_workbook

book = Workbook()
sheet = book.active
sheet.title = '计算表'
sheet.append(['临时表格检查样本（虚构数据）'])
sheet.merge_cells('A1:F1')
sheet.append(['类别', '数量', '单价', '金额', 'A类合计', '备注'])
sheet.append(['A', 2, 10, '=B3*C3', '=SUMIF(A3:A5,"A",D3:D5)', '原备注'])
sheet.append(['B', 3, 20, '=B4*C4'])
sheet.append(['A', 1, 5, '=B5*C5', '=SUM(D3:D5)'])
sheet['F5'] = '全类别合计见左侧'
sheet['F3'].comment = Comment('保留原批注', '测试作者')
validation = DataValidation(type='list', formula1='"A,B"')
sheet.add_data_validation(validation)
validation.add('A3:A5')
for row in sheet:
    for cell in row:
        cell.font = Font(name='Microsoft YaHei', size=12)
        cell.alignment = Alignment(vertical='center', wrap_text=True)
for row_num in (1, 2):
    sheet.row_dimensions[row_num].height = 29
    for cell in sheet[row_num]:
        cell.font = Font(name='Microsoft YaHei', size=14, bold=True)
        cell.fill = PatternFill('solid', fgColor='E5EDF2')
for col in 'ABCDEF':
    sheet.column_dimensions[col].width = 18 if col != 'F' else 27
for row_num in (3, 4, 5):
    sheet.row_dimensions[row_num].height = 30
    for col in 'CDE':
        sheet[f'{col}{row_num}'].number_format = '0.00'
sheet.print_area = 'A1:F5'
sheet.page_setup.orientation = 'landscape'
sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
sheet.page_setup.fitToWidth = 1
sheet.page_setup.fitToHeight = 1
sheet.sheet_properties.pageSetUpPr.fitToPage = True
raw = fixtures / 'missing-cache.xlsx'
book.save(raw)
recalculated = root / 'baseline-recalculated.xlsx'
receipt = convert('recalc', raw, recalculated, find_soffice(None), 120)
(root / 'baseline-recalc.json').write_text(json.dumps(receipt, ensure_ascii=False, indent=2), encoding='utf-8')
source = fixtures / 'calculated.xlsx'
merge_caches(raw, recalculated, source)
report = inspect_workbook(source)
assert report['formula_count'] == report['formula_cache_count'] == 5
assert report['formula_error_count'] == 0
cells = report['sheets']['计算表']['cells']
assert float(cells['E3']['cached']) == 25
assert float(cells['E5']['cached']) == 85
values = Workbook()
vs = values.active
vs.title = '清单'
vs.append(['编号', '名称', '数量'])
vs.append(['001', '虚构模块', 2])
vs.append(['002', '虚构线缆', 3])
values.save(fixtures / 'values.xlsx')
(root / 'fixture-hashes.json').write_text(json.dumps({p.name: hashlib.sha256(p.read_bytes()).hexdigest()
    for p in fixtures.iterdir()}, indent=2), encoding='utf-8')
print(json.dumps({'status': 'PASS', 'formulas': 5, 'cached': 5, 'sum_A': 25, 'total': 85}))
