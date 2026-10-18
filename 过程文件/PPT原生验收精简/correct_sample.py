from pathlib import Path
from pptx import Presentation
from pptx.util import Pt

root = Path(__file__).resolve().parent
target = root / 'sample_v03.pptx'
if target.exists():
    raise FileExistsError(target)
prs = Presentation(root / 'sample.pptx')
prs.slide_width = 12192000
for slide in prs.slides:
    title = slide.shapes[0]
    title.name = 'Title'
    for paragraph in title.text_frame.paragraphs:
        paragraph.font.size = Pt(36)
prs.save(target)
print(target)
