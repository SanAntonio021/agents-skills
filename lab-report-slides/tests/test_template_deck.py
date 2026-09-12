"""Offline regression checks for native, project-local report templates."""
import copy
import hashlib
import json
from pathlib import Path
import sys
import tempfile
import unittest
import zipfile

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import render_deck
import template_deck


class TemplateDeckTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.source = self.root / "source.pptx"
        self.output = self.root / "generated.pptx"
        self.spec_path = self.root / "template.json"
        self.old_image = self.root / "old.png"
        self.new_image = self.root / "new.png"
        Image.new("RGB", (100, 60), "red").save(self.old_image)
        Image.new("RGB", (120, 180), "blue").save(self.new_image)
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(12), Inches(8)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(230, 235, 240)
        logo = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(10), 0, Inches(1), Inches(.4))
        logo.fill.solid()
        logo.fill.fore_color.rgb = RGBColor(10, 20, 30)
        logo.text = "LAB LOGO"
        self.logo_id = logo.shape_id
        slots = {}
        for index, role in enumerate(("title", "date", "summary", "body")):
            shape = slide.shapes.add_textbox(Inches(.5), Inches(.5 + index * .65), Inches(8), Inches(.55))
            shape.text = "OLD_PRIVATE_" + role
            shape.text_frame.paragraphs[0].runs[0].font.size = Pt(24)
            slots[role] = shape.shape_id
        slots["images"], slots["captions"] = [], []
        for index in range(2):
            pic = slide.shapes.add_picture(str(self.old_image), Inches(.5 + index * 5), Inches(3.5), Inches(4), Inches(2.4))
            slots["images"].append(pic.shape_id)
            caption = slide.shapes.add_textbox(pic.left, Inches(6), pic.width, Inches(.5))
            caption.text = "OLD_PRIVATE_CAPTION_" + str(index)
            slots["captions"].append(caption.shape_id)
        slide.notes_slide.notes_text_frame.text = "OLD_PRIVATE_NOTES"
        unused = prs.slides.add_slide(prs.slide_layouts[6])
        unused.shapes.add_textbox(0, 0, Inches(5), Inches(1)).text = "OLD_PRIVATE_UNUSED_PAGE"
        unused.notes_slide.notes_text_frame.text = "OLD_PRIVATE_UNUSED_NOTES"
        prs.save(self.source)
        self.source_hash = self.digest(self.source)
        self.spec = {"template_path": str(self.source), "layouts": {"result": {
            "slide": 1, "keep_shape_ids": [self.logo_id], "slots": slots}}}
        self.deck = {"title": "NEW_REPORT", "date": "20260910", "slides": [{
            "type": "result", "title": "NEW_TITLE", "summary": "NEW_SUMMARY",
            "blocks": [{"type": "text", "text": "NEW_BODY"},
                       {"type": "image", "path": str(self.new_image), "caption": "NEW_CAPTION"}]}]}

    @staticmethod
    def digest(path):
        return hashlib.sha256(path.read_bytes()).hexdigest()

    def make(self, spec=None, deck=None):
        self.spec_path.write_text(json.dumps(spec or self.spec), encoding="utf-8")
        deck = deck or self.deck
        return template_deck.make_template_pptx(
            deck, render_deck.validate_deck(deck, self.root), self.output, self.spec_path, template_test=True)

    def test_preserves_size_native_logo_background_and_source(self):
        assets = self.make()
        self.assertEqual(self.digest(self.source), self.source_hash)
        prs = Presentation(self.output)
        self.assertEqual((prs.slide_width, prs.slide_height), (Inches(12), Inches(8)))
        self.assertEqual(len(prs.slides), 1)
        slide = prs.slides[0]
        logo = next(s for s in slide.shapes if s.shape_id == self.logo_id)
        self.assertEqual(logo.text, "LAB LOGO")
        self.assertEqual(logo.shape_type, 1)
        self.assertEqual(logo.fill.fore_color.rgb, RGBColor(10, 20, 30))
        self.assertEqual(slide.background.fill.fore_color.rgb, RGBColor(230, 235, 240))
        text = "\n".join(s.text for s in slide.shapes if s.has_text_frame)
        for value in ("NEW_TITLE", "NEW_SUMMARY", "NEW_BODY", "NEW_CAPTION", "20260910"):
            self.assertIn(value, text)
        self.assertEqual(len(assets), 1)
        self.assertEqual(Path(assets[0]["path"]), self.new_image)

    def test_project_topic_title_preserves_template_size_and_separate_styles(self):
        deck = copy.deepcopy(self.deck)
        deck["slides"][0].update(project="双向通信项目", title="功率扫描")
        self.make(deck=deck)
        title_id = self.spec["layouts"]["result"]["slots"]["title"]
        title = next(s for s in Presentation(self.output).slides[0].shapes if s.shape_id == title_id)
        self.assertEqual(title.text, "双向通信项目   功率扫描")
        self.assertEqual([(r.text, r.font.size.pt, r.font.bold) for r in title.text_frame.paragraphs[0].runs],
                         [("双向通信项目", 24, True), ("   功率扫描", 20, False)])
        self.assertEqual(len(title.text_frame.paragraphs), 1)
        self.assertFalse(title.text_frame.word_wrap)
        self.assertEqual(self.digest(self.source), self.source_hash)

    def test_project_title_requires_an_explicit_template_size(self):
        prs = Presentation(self.source)
        title_id = self.spec["layouts"]["result"]["slots"]["title"]
        title = next(s for s in prs.slides[0].shapes if s.shape_id == title_id)
        title.text_frame.paragraphs[0].runs[0].font.size = None
        prs.save(self.source)
        deck = copy.deepcopy(self.deck)
        deck["slides"][0]["project"] = "Project"
        with self.assertRaisesRegex(ValueError, "explicit title font size"):
            self.make(deck=deck)
        self.assertFalse(self.output.exists())

    def test_removes_old_private_text_notes_other_slides_and_images(self):
        self.make()
        with zipfile.ZipFile(self.output) as package:
            for name in package.namelist():
                if name.endswith(".xml"):
                    self.assertNotIn(b"OLD_PRIVATE", package.read(name), name)
            old = self.old_image.read_bytes()
            self.assertFalse(any(package.read(n) == old for n in package.namelist() if n.startswith("ppt/media/")))

    def test_fewer_images_remove_unused_picture_and_caption(self):
        self.make()
        shapes = list(Presentation(self.output).slides[0].shapes)
        pictures = [s for s in shapes if s.shape_type == 13]
        self.assertEqual(len(pictures), 1)
        ids = {s.shape_id for s in shapes}
        slots = self.spec["layouts"]["result"]["slots"]
        self.assertNotIn(slots["images"][1], ids)
        self.assertNotIn(slots["captions"][1], ids)
        # A portrait result must fit the landscape slot without distortion.
        self.assertAlmostEqual(pictures[0].width / pictures[0].height, 120 / 180, places=4)

    def test_too_many_images_fail_without_output(self):
        deck = copy.deepcopy(self.deck)
        deck["slides"][0]["blocks"].extend([copy.deepcopy(deck["slides"][0]["blocks"][1]) for _ in range(2)])
        with self.assertRaises(ValueError):
            self.make(deck=deck)
        self.assertFalse(self.output.exists())

    def test_invalid_slot_ids_and_types_fail_without_output(self):
        for role, value in (("title", 9999), ("title", self.spec["layouts"]["result"]["slots"]["images"][0]),
                            ("images", [self.spec["layouts"]["result"]["slots"]["title"]]),
                            ("summary", self.spec["layouts"]["result"]["slots"]["title"])):
            with self.subTest(role=role, value=value):
                spec = copy.deepcopy(self.spec)
                spec["layouts"]["result"]["slots"][role] = value
                with self.assertRaises(ValueError):
                    self.make(spec=spec)
                self.assertFalse(self.output.exists())

    def test_existing_output_is_never_overwritten(self):
        self.output.write_bytes(b"existing-user-file")
        with self.assertRaises((ValueError, FileExistsError)):
            self.make()
        self.assertEqual(self.output.read_bytes(), b"existing-user-file")

    def test_relative_template_path_uses_spec_directory(self):
        spec = copy.deepcopy(self.spec)
        spec["template_path"] = self.source.name
        self.make(spec=spec)
        self.assertEqual(len(Presentation(self.output).slides), 1)

    def test_potx_becomes_standard_pptx_without_mutating_template(self):
        potx = self.root / "source.potx"
        normal = b"application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"
        template = b"application/vnd.openxmlformats-officedocument.presentationml.template.main+xml"
        with zipfile.ZipFile(self.source) as source, zipfile.ZipFile(potx, "w") as target:
            for item in source.infolist():
                data = source.read(item.filename)
                target.writestr(item, data.replace(normal, template) if item.filename == "[Content_Types].xml" else data)
        before = self.digest(potx)
        spec = copy.deepcopy(self.spec)
        spec["template_path"] = str(potx)
        self.make(spec=spec)
        self.assertEqual(self.digest(potx), before)
        self.assertEqual(len(Presentation(self.output).slides), 1)
        with zipfile.ZipFile(self.output) as package:
            self.assertIn(normal, package.read("[Content_Types].xml"))
            self.assertNotIn(template, package.read("[Content_Types].xml"))

    def test_external_relationship_is_rejected(self):
        prs = Presentation(self.source)
        shape = next(s for s in prs.slides[0].shapes if s.shape_id == self.logo_id)
        shape.text_frame.paragraphs[0].runs[0].hyperlink.address = "https://example.invalid/"
        prs.save(self.source)
        with self.assertRaises(ValueError):
            self.make()
        self.assertFalse(self.output.exists())

    def test_inspection_is_json_serializable_and_reports_slot_ids(self):
        report = template_deck.inspect_template(self.source)
        self.assertEqual(json.loads(json.dumps(report))["slide_count"], 2)
        self.assertEqual(report["size_emu"], [Inches(12), Inches(8)])
        shapes = report["slides"][0]["shapes"]
        title_id = self.spec["layouts"]["result"]["slots"]["title"]
        title = next(shape for shape in shapes if shape["shape_id"] == title_id)
        self.assertIn("OLD_PRIVATE_title", title["text"])
        self.assertEqual(len(title["bounds_emu"]), 4)
        self.assertEqual(self.digest(self.source), self.source_hash)

    def test_master_picture_and_theme_remain_reachable(self):
        master_image = self.root / "master-logo.png"
        Image.new("RGB", (80, 80), "green").save(master_image)
        prs = Presentation(self.source)
        pic = next(s for s in prs.slides[0].shapes if s.shape_type == 13)
        native = copy.deepcopy(pic._element)
        master = prs.slide_master
        _, image_rel = master.part.get_or_add_image_part(str(master_image))
        native.find(".//" + qn("a:blip")).set(qn("r:embed"), image_rel)
        props = native.find(".//" + qn("p:cNvPr"))
        props.set("id", str(max(s.shape_id for s in master.shapes) + 1))
        props.set("name", "MASTER_DESIGN_LOGO")
        master.shapes._spTree.insert_element_before(native, "p:extLst")
        prs.save(self.source)
        self.make()
        result = Presentation(self.output)
        master_pictures = [s for s in result.slide_master.shapes if s.shape_type == 13]
        self.assertTrue(any(s.image.blob == master_image.read_bytes() for s in master_pictures))
        with zipfile.ZipFile(self.output) as package:
            self.assertTrue(any(n.startswith("ppt/theme/") and n.endswith(".xml") for n in package.namelist()))

    def test_shape_override_changes_only_the_requested_native_box(self):
        spec = copy.deepcopy(self.spec)
        title_id = spec["layouts"]["result"]["slots"]["title"]
        box = [Inches(1), Inches(.2), Inches(9), Inches(.7)]
        spec["layouts"]["result"]["shape_overrides"] = {str(title_id): {"box_emu": box}}
        self.make(spec=spec)
        title = next(s for s in Presentation(self.output).slides[0].shapes if s.shape_id == title_id)
        self.assertEqual([title.left, title.top, title.width, title.height], box)
        self.assertEqual(title.text_frame.paragraphs[0].runs[0].font.size, Pt(24))
        self.assertEqual(self.digest(self.source), self.source_hash)

    def test_invalid_shape_override_fails_without_output(self):
        for box in ([0, 0, -1, 100], [0, 0, Inches(13), 100], [0, 0, 100], [0, 0, True, 100]):
            with self.subTest(box=box):
                spec = copy.deepcopy(self.spec)
                spec["layouts"]["result"]["shape_overrides"] = {str(self.logo_id): {"box_emu": box}}
                with self.assertRaises(ValueError):
                    self.make(spec=spec)
                self.assertFalse(self.output.exists())

    def test_next_steps_require_an_explicit_layout(self):
        deck = copy.deepcopy(self.deck)
        deck["slides"].append({"type": "next_steps", "title": "Next", "next_steps": ["First action", "Second action"]})
        with self.assertRaises(ValueError):
            self.make(deck=deck)
        self.assertFalse(self.output.exists())
        spec = copy.deepcopy(self.spec)
        spec["layouts"]["next_steps"] = copy.deepcopy(spec["layouts"]["result"])
        self.make(spec=spec, deck=deck)
        prs = Presentation(self.output)
        self.assertEqual(len(prs.slides), 2)
        shapes = list(prs.slides[1].shapes)
        self.assertEqual(sum(s.shape_type == 13 for s in shapes), 0)
        self.assertTrue(any(s.has_text_frame and s.text == "1. First action\n2. Second action" for s in shapes))


if __name__ == "__main__":
    unittest.main()
