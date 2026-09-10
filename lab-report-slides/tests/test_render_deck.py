import json
import sys
import tempfile
import unittest
import zipfile
import shutil
from pathlib import Path

from PIL import Image
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import render_deck  # noqa: E402


class RenderDeckTests(unittest.TestCase):
    def test_exports_actual_pptx_renders_and_preserves_editable_objects(self):
        with tempfile.TemporaryDirectory() as temp:
            base = Path(temp)
            image = base / "result.png"
            Image.new("RGB", (320, 180), "#4472C4").save(image)
            deck = base / "deck.json"
            deck.write_text(json.dumps({
                "title": "测试汇报",
                "date": "20260715",
                "slides": [{"title": "测试结果", "status": "已完成", "blocks": [{"type": "text", "text": "完成测试"}, {"type": "image", "path": str(image), "caption": "结果图"}]}]
            }, ensure_ascii=False), encoding="utf-8")
            output = base / "out"
            manifest = render_deck.render(deck, output, "20260715")

            self.assertEqual(manifest["slide_count"], 1)
            self.assertTrue((output / "20260715.html").stat().st_size > 0)
            self.assertTrue((output / "20260715.pdf").stat().st_size > 0)
            self.assertTrue(zipfile.is_zipfile(output / "20260715.pptx"))
            presentation = Presentation(output / "20260715.pptx")
            slide = presentation.slides[0]
            self.assertTrue(any(shape.has_text_frame and "完成测试" in shape.text for shape in slide.shapes))
            pictures = [shape for shape in slide.shapes if shape.shape_type == 13]
            self.assertEqual(1, len(pictures))
            self.assertLess(pictures[0].height, presentation.slide_height * 0.8)
            self.assertEqual(manifest["render_source"], "pptx")
            self.assertTrue(manifest["editable_text"])
            self.assertEqual(manifest["assets"][0]["path"], str(image.resolve()))
            with Image.open(output / "20260715_01.png") as rendered:
                self.assertEqual(rendered.size, (1600, 900))
            second = render_deck.render(deck, output, "20260715")
            self.assertEqual(second["stem"], "20260715_v2")
            self.assertTrue((output / "20260715_v2.pptx").exists())

    def test_summary_body_and_numbered_next_steps_use_separate_editable_regions(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            image = root / "result.png"
            Image.new("RGB", (800, 400), "white").save(image)
            deck = {"date": "20260909", "slides": [
                {"title": "多流系统", "summary": "完成三组软件译码。", "font_sizes": {"title": 31, "body": 19},
                 "blocks": [{"type": "text", "text": "三组均同步通过"}, {"type": "image", "path": str(image)}]},
                {"title": "下一步工作", "type": "next_steps", "next_steps": ["开展实测对比", "补充结果记录"]}]}
            render_deck.make_pptx(deck, render_deck.validate_deck(deck, root), root / "out.pptx")
            slides = Presentation(root / "out.pptx").slides
            texts = [shape for shape in slides[0].shapes if shape.has_text_frame and shape.text]
            self.assertEqual(sum(shape.text == "完成三组软件译码。" for shape in texts), 1)
            self.assertTrue(all("20260909" not in shape.text for shape in texts))
            summary = next(shape for shape in texts if shape.text == "完成三组软件译码。")
            body = next(shape for shape in texts if shape.text == "三组均同步通过")
            picture = next(shape for shape in slides[0].shapes if shape.shape_type == 13)
            self.assertLess(summary.top, body.top)
            self.assertLessEqual(body.left + body.width, picture.left)
            self.assertEqual(body.text_frame.paragraphs[0].runs[0].font.size.pt, 19)
            title = next(shape for shape in texts if shape.text == "多流系统")
            self.assertEqual(title.text_frame.paragraphs[0].runs[0].font.size.pt, 31)
            numbered = [shape for shape in slides[1].shapes if shape.has_text_frame and shape.text.startswith("1.")]
            self.assertEqual(len(numbered), 1)
            self.assertEqual(numbered[0].text, "1. 开展实测对比\n2. 补充结果记录")

    def test_next_steps_cannot_be_intermediate_or_mixed_with_body(self):
        for slides in [
            [{"type": "next_steps", "next_steps": ["Action"]}, {"body": "Result"}],
            [{"type": "next_steps", "next_steps": []}],
            [{"type": "next_steps", "next_steps": ["Action"], "body": "Extra"}],
        ]:
            with self.assertRaises(ValueError):
                render_deck.validate_deck({"allow_text_only": True, "slides": slides}, Path.cwd())

    def test_missing_image_is_rejected_instead_of_creating_a_placeholder(self):
        with tempfile.TemporaryDirectory() as temp:
            with self.assertRaisesRegex(ValueError, "Missing research image"):
                render_deck.validate_deck({"slides": [{"title": "Result", "blocks": [{"type": "image", "path": "missing.png"}]}]}, Path(temp))

    def test_profile_controls_layout_and_six_images_are_preserved(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            picture = root / "photo.png"
            Image.new("RGB", (300, 200), "white").save(picture)
            profile = json.loads(render_deck.PROFILE_PATH.read_text(encoding="utf-8"))
            profile["layout"]["title"].update(box=[0.3, 0.04, 12.5, 0.68], color="123456", font_size=36)
            profile_path = root / "profile.json"
            profile_path.write_text(json.dumps(profile), encoding="utf-8")
            deck = {"profile_path": str(profile_path), "slides": [{"section": "Chapter", "title": "Experiment",
                "layout": "wide-strip", "summary": "Result", "blocks": [{"type": "image", "path": str(picture)} for _ in range(6)]}]}
            render_deck.make_pptx(deck, render_deck.validate_deck(deck, root), root / "out.pptx")
            slide = Presentation(root / "out.pptx").slides[0]
            title = next(s for s in slide.shapes if s.has_text_frame and s.text == "Chapter")
            self.assertAlmostEqual(title.left / 914400, 0.3)
            run = title.text_frame.paragraphs[0].runs[0]
            self.assertEqual(run.font.size.pt, 36)
            self.assertEqual(str(run.font.color.rgb), "123456")
            pictures = [s for s in slide.shapes if s.shape_type == 13]
            self.assertEqual(len(pictures), 6)
            self.assertTrue(all(p.top > pictures[0].top for p in pictures[1:]))
            boxes = render_deck.image_boxes(3, [0, 0, 12, 5], "stacked-left")
            self.assertGreater(boxes[2][3], boxes[0][3])
            self.assertEqual(render_deck.image_boxes(0, [0, 0, 12, 5], None), [])

    def test_text_only_requires_an_explicit_choice(self):
        deck = {"slides": [{"title": "Summary", "body": "Text"}]}
        with self.assertRaisesRegex(ValueError, "No experimental plots"):
            render_deck.validate_deck(deck, Path.cwd())
        deck["allow_text_only"] = True
        self.assertEqual(len(render_deck.validate_deck(deck, Path.cwd())), 1)

    def test_empty_material_and_results_without_images_are_rejected(self):
        with self.assertRaises(ValueError):
            render_deck.validate_deck({"slides": []}, Path.cwd())
        with self.assertRaises(ValueError):
            render_deck.validate_deck({"allow_text_only": True, "slides": [{"type": "result", "body": "Result"}]}, Path.cwd())

    def test_negative_measurements_are_not_stripped_as_bullets(self):
        rows = render_deck.text_rows({"type": "text", "text": "-44 dBm\n- condition\n* another condition"})
        self.assertEqual(rows, [("-44 dBm", False), ("condition", False), ("another condition", False)])

    def test_font_fit_uses_standard_steps_and_preserves_explicit_small_size(self):
        self.assertEqual(render_deck.text_size([("Title", True)], 12, 0.7, 31), 31)
        self.assertEqual(render_deck.text_size([("Title", True)], 12, 0.5, 31), 24)
        self.assertEqual(render_deck.text_size([("Caption", False)], 5, 0.4, 17), 17)
        with self.assertRaisesRegex(ValueError, "does not fit"):
            render_deck.text_size([("Caption", False)], 5, 0.2, 17)

    def test_long_text_requires_rewriting_instead_of_silent_overflow(self):
        with self.assertRaisesRegex(ValueError, "does not fit"):
            render_deck.text_size([("A" * 5000, False)], 3.6, 4.9, 21)

    def test_base_name_cannot_escape_output_directory(self):
        with self.assertRaises(ValueError):
            render_deck.choose_stem(Path.cwd(), "../report")

    def test_active_output_reservation_prevents_name_collisions(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            with render_deck.reserve_stem(root, "report") as first:
                with render_deck.reserve_stem(root, "report") as second:
                    self.assertEqual(first, "report")
                    self.assertEqual(second, "report_v2")
            self.assertFalse(list(root.glob("*.reserve")))

    def test_partial_manifest_and_png_outputs_are_not_overwritten(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            manifest = root / "report.manifest.json"
            manifest.write_text("original", encoding="utf-8")
            self.assertEqual(render_deck.choose_stem(root, "report"), "report_v2")
            (root / "report_v2_01.png").write_bytes(b"retained")
            self.assertEqual(render_deck.choose_stem(root, "report"), "report_v3")
            self.assertEqual(manifest.read_text(encoding="utf-8"), "original")

    def test_svg_is_rendered_with_existing_image_runtime(self):
        sharp = Path.home() / ".cache/codex-runtimes/codex-primary-runtime/dependencies/node/node_modules/sharp"
        if not shutil.which("node") or not sharp.is_dir():
            self.skipTest("Existing Node/sharp runtime is not available")
        with tempfile.TemporaryDirectory() as temp:
            source = Path(temp) / "plot.svg"
            original = b'<svg xmlns="http://www.w3.org/2000/svg" width="200" height="120"><rect width="200" height="120" fill="white"/><path d="M10 100L90 60L180 20" stroke="blue" fill="none"/></svg>'
            source.write_bytes(original)
            raster = render_deck.raster_bytes(source)
            self.assertTrue(raster.startswith(b"\x89PNG"))
            self.assertEqual(original, source.read_bytes())


if __name__ == "__main__":
    unittest.main()
