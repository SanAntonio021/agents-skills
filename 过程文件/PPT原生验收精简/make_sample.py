from pathlib import Path
import json
import hashlib
import math
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt

root = Path(__file__).resolve().parent
sample = root / 'sample.pptx'
if sample.exists():
    raise FileExistsError(sample)
x = [i / 100 for i in range(101)]
fig, ax = plt.subplots(figsize=(7, 3))
ax.plot(x, [math.sin(2 * math.pi * v) for v in x], color='#176C8E')
ax.set(xlabel='Time (s)', ylabel='Amplitude (a.u.)', title='Synthetic test signal')
fig.tight_layout()
chart = root / 'sample-wave.png'
fig.savefig(chart, dpi=160)
plt.close(fig)
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)

def text(slide, label, x, y, w, h, size=22):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    shape.text = label
    for p in shape.text_frame.paragraphs:
        p.font.name = 'Microsoft YaHei'
        p.font.size = Pt(size)
        p.font.color.rgb = RGBColor.from_string('203040')
    return shape

s = prs.slides.add_slide(prs.slide_layouts[6])
text(s, 'PowerPoint 原生验收样例', .65, .45, 12, .8, 32)
text(s, '临时合成样本：中文、图片与单位标注', .65, 1.45, 12, .7)
s.shapes.add_picture(str(chart), Inches(.85), Inches(2.45), width=Inches(8.4))
text(s, '检查内容\n文字未裁切\n图片完整\n源文件保持不变', 9.55, 2.8, 3.0, 2.4, 20)
s = prs.slides.add_slide(prs.slide_layouts[6])
text(s, '原生表格与后续编辑', .65, .45, 12, .8, 32)
text(s, '以下数值只用于校验排版，均为虚构测试数据。', .65, 1.45, 12, .7)
table = s.shapes.add_table(4, 3, Inches(.85), Inches(2.65), Inches(11.5), Inches(2.45)).table
for i, row in enumerate([['方案', 'Channel', '设定功率 (dBm)'], ['方案 A', '1', '-10.00'],
                          ['方案 A', '2', '-8.50'], ['方案 B', '1', '-12.00']]):
    for j, value in enumerate(row):
        cell = table.cell(i, j)
        cell.text = value
        for p in cell.text_frame.paragraphs:
            p.font.name = 'Microsoft YaHei'
            p.font.size = Pt(22)
text(s, '交付检查：结构 → PowerPoint 打开与导出 → 逐页查看', .85, 5.65, 11.5, .7, 20)
prs.save(sample)
(root / 'sample-source.json').write_text(json.dumps({'file': str(sample),
    'sha256': hashlib.sha256(sample.read_bytes()).hexdigest(), 'slides': 2}, indent=2), encoding='utf-8')
print(sample)
